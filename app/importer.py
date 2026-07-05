"""
Formula import — parse Excel / CSV / pasted text / PDF into component rows.

Every parser returns the same shape:

    {
      'components': [{'name': str, 'cas': str, 'pct': float}],
      'warnings':   [str],
      'source':     'xlsx' | 'csv' | 'text' | 'pdf',
    }

Parsed data is always shown to the user in the review grid before an SDS is
generated — this module never writes to any database and never generates a
document on its own. That review step is the safety valve for imperfect
inputs (especially PDFs).
"""

import csv
import io
import re

# CAS: 2-7 digits, 2 digits, 1 check digit
_CAS_RE = re.compile(r'\b(\d{2,7}-\d{2}-\d)\b')

# Header cell → column role
_NAME_HEADERS = {'name', 'chemical', 'chemical name', 'ingredient', 'component',
                 'substance', 'material', 'raw material', 'description', 'product'}
_CAS_HEADERS = {'cas', 'cas no', 'cas no.', 'cas#', 'cas #', 'casno', 'cas number',
                'cas-nr', 'cas rn', 'cas-rn', 'cas registry'}
_PCT_HEADERS = {'%', 'percent', 'pct', 'wt', 'wt%', 'wt %', 'weight', 'weight %',
                'weight percent', 'conc', 'conc.', 'concentration', 'amount',
                'w/w', '% w/w', '%w/w', 'wgt'}


# ── Public API ────────────────────────────────────────────────────────────────

def parse_upload(file_storage):
    """Parse an uploaded file (werkzeug FileStorage) by extension."""
    filename = (file_storage.filename or '').lower()
    raw = file_storage.read()

    if filename.endswith('.xlsx'):
        return _parse_xlsx(raw)
    if filename.endswith('.xls'):
        return _empty('xlsx', ['Old .xls format is not supported — '
                               'please re-save as .xlsx and try again.'])
    if filename.endswith('.csv'):
        return _parse_csv(raw)
    if filename.endswith('.pdf'):
        return _parse_pdf(raw)
    if filename.endswith(('.txt', '.tsv')):
        return _parse_text(raw.decode('utf-8', errors='replace'))
    return _empty('text', [f'Unsupported file type: {filename or "unknown"}. '
                           'Use .xlsx, .csv, .txt, or .pdf.'])


def parse_text(text):
    """Parse pasted text (spreadsheet rows, tab/comma/space separated)."""
    return _parse_text(text)


# ── Format parsers ────────────────────────────────────────────────────────────

def _parse_xlsx(raw):
    try:
        from openpyxl import load_workbook
    except ImportError:
        return _empty('xlsx', ['openpyxl is not installed — run: pip install openpyxl'])
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        return _empty('xlsx', [f'Could not read the Excel file: {exc}'])
    ws = wb.active
    rows = []
    for r in ws.iter_rows(values_only=True):
        rows.append(['' if c is None else str(c) for c in r])
    wb.close()
    return _rows_to_result(rows, 'xlsx')


def _parse_csv(raw):
    text = raw.decode('utf-8', errors='replace') if isinstance(raw, bytes) else raw
    sample = text[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=',;\t|')
    except csv.Error:
        dialect = csv.excel
    rows = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
    return _rows_to_result(rows, 'csv')


def _parse_text(text):
    rows = []
    for line in text.splitlines():
        if not line.strip():
            continue
        rows.append(_split_line(line))
    return _rows_to_result(rows, 'text')


def _parse_pdf(raw):
    try:
        from pypdf import PdfReader
    except BaseException:   # ImportError, or a broken native crypto backend in some envs
        return _empty('pdf', ['PDF support is unavailable — run: pip install pypdf'])
    try:
        reader = PdfReader(io.BytesIO(raw))
        text = '\n'.join((page.extract_text() or '') for page in reader.pages)
    except Exception as exc:
        return _empty('pdf', [f'Could not read the PDF: {exc}'])

    if not text.strip():
        return _empty('pdf', ['No text found in the PDF. Scanned/image PDFs are not '
                              'supported (no OCR). Enter the components manually.'])

    # A composition table row typically has a CAS number and a percentage.
    components = []
    warnings = []
    for line in text.splitlines():
        if not _CAS_RE.search(line):
            continue
        cas = _CAS_RE.search(line).group(1)
        pct = _find_percent(line)
        name = _name_from_line(line, cas)
        if name or pct is not None:
            components.append({'name': name, 'cas': cas, 'pct': pct or 0.0})

    if not components:
        warnings.append('No composition table (CAS + percent rows) was found in the PDF. '
                        'The layout may not be machine-readable — enter components manually.')
    else:
        warnings.append('Imported from PDF text — verify every row against the source '
                        'document, as PDF layouts can scramble during extraction.')
    return {'components': components, 'warnings': warnings, 'source': 'pdf'}


# ── Row-based mapping (xlsx / csv / text share this) ─────────────────────────

def _rows_to_result(rows, source):
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return _empty(source, ['No data rows found.'])

    warnings = []
    header_idx, mapping = _detect_header(rows)

    if mapping is None:
        # No header detected — infer columns positionally
        mapping = _infer_columns(rows)
        data_rows = rows
        warnings.append('No header row detected — columns were guessed from the data. '
                        'Please check each row.')
    else:
        data_rows = rows[header_idx + 1:]

    if 'name' not in mapping and 'cas' not in mapping:
        return _empty(source, ['Could not identify a chemical-name or CAS column. '
                               'Expected columns like "Chemical", "CAS", and "%".'])

    components = []
    for r in data_rows:
        name = _cell(r, mapping.get('name'))
        cas = _cell(r, mapping.get('cas'))
        pct_raw = _cell(r, mapping.get('pct'))

        # Recover a CAS embedded in the name cell
        if not cas and name:
            m = _CAS_RE.search(name)
            if m:
                cas = m.group(1)
                name = name.replace(cas, '').strip(' ,;')

        cas = cas if _valid_cas(cas) else (cas or '')
        pct = _to_float(pct_raw)
        if not name and not cas:
            continue
        components.append({'name': name, 'cas': cas, 'pct': pct or 0.0})

    if not components:
        return _empty(source, ['No component rows could be read from the file.'])

    if mapping.get('pct') is None:
        warnings.append('No percentage column found — enter weight percentages manually.')
    return {'components': components, 'warnings': warnings, 'source': source}


def _detect_header(rows):
    """Find a header row in the first few rows and map columns to roles."""
    for idx, row in enumerate(rows[:5]):
        mapping = {}
        for col, cell in enumerate(row):
            key = str(cell).strip().lower().rstrip(':')
            if key in _NAME_HEADERS and 'name' not in mapping:
                mapping['name'] = col
            elif key in _CAS_HEADERS and 'cas' not in mapping:
                mapping['cas'] = col
            elif (key in _PCT_HEADERS or '%' in key) and 'pct' not in mapping:
                mapping['pct'] = col
        if 'name' in mapping or 'cas' in mapping:
            return idx, mapping
    return -1, None


def _infer_columns(rows):
    """Guess columns positionally when there's no header."""
    ncols = max(len(r) for r in rows)
    mapping = {}

    # CAS column: the one with the most valid CAS numbers
    best_cas, best_cas_hits = None, 0
    for col in range(ncols):
        hits = sum(1 for r in rows if _valid_cas(_cell(r, col)))
        if hits > best_cas_hits:
            best_cas, best_cas_hits = col, hits
    if best_cas is not None and best_cas_hits:
        mapping['cas'] = best_cas

    # Percent column: numeric values mostly within 0-100
    best_pct, best_pct_hits = None, 0
    for col in range(ncols):
        if col == mapping.get('cas'):
            continue
        hits = 0
        for r in rows:
            v = _to_float(_cell(r, col))
            if v is not None and 0 < v <= 100:
                hits += 1
        if hits > best_pct_hits:
            best_pct, best_pct_hits = col, hits
    if best_pct is not None and best_pct_hits:
        mapping['pct'] = best_pct

    # Name column: the remaining column with the longest average text
    best_name, best_len = None, 0
    for col in range(ncols):
        if col in (mapping.get('cas'), mapping.get('pct')):
            continue
        avg = _avg_text_len(rows, col)
        if avg > best_len:
            best_name, best_len = col, avg
    if best_name is not None:
        mapping['name'] = best_name
    return mapping


# ── Small helpers ─────────────────────────────────────────────────────────────

def _split_line(line):
    if '\t' in line:
        return [c.strip() for c in line.split('\t')]
    if ',' in line:
        return [c.strip() for c in line.split(',')]
    return [c for c in re.split(r'\s{2,}', line.strip()) if c]


def _cell(row, col):
    if col is None or col >= len(row):
        return ''
    return str(row[col]).strip()


def _avg_text_len(rows, col):
    vals = [len(_cell(r, col)) for r in rows if _cell(r, col)
            and not _valid_cas(_cell(r, col)) and _to_float(_cell(r, col)) is None]
    return sum(vals) / len(vals) if vals else 0


def _find_percent(line):
    m = re.search(r'(\d+(?:\.\d+)?)\s*%', line)
    if m:
        return float(m.group(1))
    return None


def _name_from_line(line, cas):
    # Text before the CAS number, minus leading numbering
    before = line.split(cas)[0].strip()
    before = re.sub(r'^\s*\d+[.)]\s*', '', before)
    before = re.sub(r'\s+', ' ', before).strip(' ,;|')
    return before


def _to_float(value):
    if value is None:
        return None
    s = str(value).strip().replace('%', '').replace(',', '.')
    m = re.search(r'-?\d+(?:\.\d+)?', s)
    if not m:
        return None
    try:
        return float(m.group(0))
    except ValueError:
        return None


def _valid_cas(s):
    s = str(s).strip()
    if not re.fullmatch(r'\d{2,7}-\d{2}-\d', s):
        return False
    digits = s.replace('-', '')
    body, check = digits[:-1], int(digits[-1])
    return sum(int(d) * i for i, d in enumerate(reversed(body), start=1)) % 10 == check


def _empty(source, warnings):
    return {'components': [], 'warnings': warnings, 'source': source}
